#!/usr/bin/env python3
"""Escribe una extracción validada en la matriz de síntesis.

Uso:
    python3 escribir_matriz.py salida.json Libro_de_codigos_extraccion_v2.json matriz_entrada.xlsx matriz_salida.xlsx [auditoria.json]

Qué hace:
- Reemplaza la fila cuyo Estudio coincide o agrega la fila al final de la hoja de extracción.
- Escribe los 54 valores en el orden de posicion del libro de códigos; null queda como celda vacía.
- Pone en cada celda con evidencia un comentario con la página, la ubicación y la cita textual.
- Colorea: amarillo = valor inferido; naranja = discrepancia en el artículo; rojo = hallazgo de auditoría no resuelto.
- Crea o actualiza las hojas Trazabilidad (una fila por dato) y Auditoria (hallazgos de los agentes auditores).
No modifica matriz_entrada.xlsx: siempre guarda en matriz_salida.xlsx.
"""
import json, sys
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

AMARILLO = PatternFill("solid", fgColor="FFF2CC")
NARANJA = PatternFill("solid", fgColor="F8CBAD")
ROJO = PatternFill("solid", fgColor="FF9999")
HOJA = "Extracción completa.csv"


def hoja_limpia(wb, nombre, encabezados):
    ws = wb[nombre] if nombre in wb.sheetnames else wb.create_sheet(nombre)
    if ws.max_row < 1 or ws.cell(1, 1).value != encabezados[0]:
        ws.delete_rows(1, ws.max_row)
        ws.append(encabezados)
    return ws


def borrar_estudio(ws, estudio):
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, 1).value == estudio:
            ws.delete_rows(r)


def main():
    if len(sys.argv) < 5:
        print(__doc__); sys.exit(2)
    salida = json.load(open(sys.argv[1], encoding="utf-8"))
    libro = json.load(open(sys.argv[2], encoding="utf-8"))
    auditoria = json.load(open(sys.argv[5], encoding="utf-8")) if len(sys.argv) > 5 else []
    wb = openpyxl.load_workbook(sys.argv[3])
    ws = wb[HOJA]
    cols = sorted(libro["columnas"].items(), key=lambda kv: kv[1]["posicion"])
    col_idx = {k: c["posicion"] for k, c in cols}

    tz = hoja_limpia(wb, "Trazabilidad", ["Estudio", "Columna", "Valor", "Página PDF", "Ubicación", "Fila tabla", "Columna tabla", "Cita textual", "Leído de figura", "Tipo de valor", "Nota"])
    au = hoja_limpia(wb, "Auditoria", ["Estudio", "Columna", "Auditor", "Veredicto", "Hallazgo", "Resolución", "Estado"])

    for row in salida:
        f, t = row["fila"], row["trazabilidad"]
        est = f["Estudio"]
        # fila destino
        destino = None
        ultima = 1
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 2).value
            if v is not None:
                ultima = r
                try:
                    if int(v) == int(est):
                        destino = r
                except (TypeError, ValueError):
                    pass
        if destino is None:
            destino = ultima + 1
        inferidas = {x["columna"].split(" (")[0] for x in t.get("inferencias", [])}
        discrep = {x["columna"] for x in t.get("discrepancias_en_el_articulo", [])}
        pendientes = {a["columna"] for a in auditoria if a.get("estudio") == est and a.get("estado") != "Resuelto"}

        for k, _ in cols:
            c = ws.cell(destino, col_idx[k])
            c.value = f[k]
            c.comment = None
            c.fill = PatternFill()
            evs = t["evidencia"].get(k, [])
            if evs:
                texto = "\n".join(
                    f"p. {e['pagina_pdf']} | {e['ubicacion']}"
                    + (f" | fila: {e['fila_tabla']}" if e.get("fila_tabla") else "")
                    + (f" | col: {e['columna_tabla']}" if e.get("columna_tabla") else "")
                    + (" | leído de figura" if e.get("leido_de_figura") else "")
                    + f"\n\"{e['cita_textual']}\"" for e in evs)
                c.comment = Comment(texto[:1500], "Extractor")
            elif f[k] is None and k in t["estado_del_dato"]:
                c.comment = Comment(t["estado_del_dato"][k], "Extractor")
            if k in pendientes:
                c.fill = ROJO
            elif k in discrep:
                c.fill = NARANJA
            elif k in inferidas:
                c.fill = AMARILLO

        borrar_estudio(tz, est)
        for k, _ in cols:
            tipo = "Faltante" if f[k] is None or f[k] in ("No indica", "No aplica", "No indica (no significativo)") else (
                "Inferido" if k in inferidas else "Literal o codificado")
            evs = t["evidencia"].get(k) or [None]
            for e in evs:
                tz.append([est, k, json.dumps(f[k], ensure_ascii=False) if not isinstance(f[k], (int, float, str)) else f[k],
                           e and e["pagina_pdf"], e and e["ubicacion"], e and e.get("fila_tabla"), e and e.get("columna_tabla"),
                           e and e["cita_textual"], e and e.get("leido_de_figura"), tipo,
                           t["estado_del_dato"].get(k) if f[k] is None else None])
        for x in t.get("inferencias", []):
            tz.append([est, x["columna"], x["valor"], None, None, None, None, None, None, "Inferencia", f"{x['regla_aplicada']}: {x['razonamiento']}"])
        for x in t.get("discrepancias_en_el_articulo", []):
            tz.append([est, x["columna"], x["valor_adoptado"], None, None, None, None, None, None, "Discrepancia del artículo", f"{x['fuente_1']} || {x['fuente_2']}"])
        for x in t.get("advertencias_pendientes", []):
            tz.append([est, None, None, None, None, None, None, None, None, "Advertencia", x])
        tz.append([est, "confianza", t["confianza"], None, None, None, None, None, None, "Resumen", f"extractor: {t['extractor']}; fecha: {t['fecha_extraccion']}; verificador: {t['verificador']}"])

        borrar_estudio(au, est)
        for a in auditoria:
            if a.get("estudio") == est:
                au.append([est, a.get("columna"), a.get("auditor"), a.get("veredicto"), a.get("hallazgo"), a.get("resolucion"), a.get("estado")])
        print(f"Estudio {est}: fila {destino} de '{HOJA}'")

    wb.save(sys.argv[4])
    print("Guardado en", sys.argv[4])


if __name__ == "__main__":
    main()
