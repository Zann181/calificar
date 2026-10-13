"""Escritores de Excel.

EscritorHeredadasOpenpyxl: vacía los valores y comentarios de las filas de datos del Excel base
(conserva formato y colores) y escribe las filas heredadas en su fila de origen.

ScriptEscritor: ejecuta recursos/escribir_matriz.py tal cual, con archivos temporales.
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.comments import Comment

from app.compartido.columnas import DefinicionDeColumna
from app.contextos.exportacion.puertos import FilaExportable, MatrizBase

MARCA_FECHA = "$fecha"


class ErrorDeEscritura(RuntimeError):
    pass


def _desde_json(valor: Any) -> Any:
    if isinstance(valor, dict) and MARCA_FECHA in valor:
        return datetime.fromisoformat(valor[MARCA_FECHA])
    return valor


class EscritorHeredadasOpenpyxl:
    def escribir(self, base: MatrizBase, filas: list[FilaExportable], columnas: list[DefinicionDeColumna]) -> bytes:
        wb = openpyxl.load_workbook(io.BytesIO(base.datos))
        ws = wb[base.hoja]
        posiciones = {c.clave: c.posicion for c in columnas}
        for r in range(2, ws.max_row + 1):
            for p in posiciones.values():
                celda = ws.cell(r, p)
                celda.value = None
                celda.comment = None

        ocupadas: set[int] = set()
        siguiente = 2
        for f in sorted(filas, key=lambda x: x.estudio):
            origen = f.fila_excel_origen
            if origen is None or origen in ocupadas:
                while siguiente in ocupadas:
                    siguiente += 1
                origen = siguiente
            r = origen
            ocupadas.add(r)
            for clave, p in posiciones.items():
                celda = ws.cell(r, p)
                celda.value = _desde_json(f.valores.get(clave))
                nota = f.notas_heredadas.get(clave)
                if nota:
                    celda.comment = Comment(nota["texto"], nota.get("autor") or "")
        salida = io.BytesIO()
        wb.save(salida)
        return salida.getvalue()


class ScriptEscritor:
    """Adaptador de EscritorDeMatriz sobre recursos/escribir_matriz.py (no se modifica el script)."""

    def __init__(self, script: Path, python: str = sys.executable, tiempo_max: int = 300) -> None:
        self._script, self._python, self._tiempo = script, python, tiempo_max

    def escribir(self, salida: list[dict[str, Any]], libro: dict[str, Any], matriz_base: bytes,
                 auditoria: list[dict[str, Any]]) -> bytes:
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "salida.json").write_text(json.dumps(salida, ensure_ascii=False), encoding="utf-8")
            (d / "libro.json").write_text(json.dumps(libro, ensure_ascii=False), encoding="utf-8")
            (d / "auditoria.json").write_text(json.dumps(auditoria, ensure_ascii=False), encoding="utf-8")
            (d / "entrada.xlsx").write_bytes(matriz_base)
            proc = subprocess.run(
                [self._python, str(self._script), "salida.json", "libro.json", "entrada.xlsx", "salida.xlsx",
                 "auditoria.json"],
                cwd=d, capture_output=True, text=True, encoding="utf-8", timeout=self._tiempo,
                env={**os.environ, "PYTHONIOENCODING": "utf-8"},
            )
            if proc.returncode != 0 or not (d / "salida.xlsx").exists():
                raise ErrorDeEscritura(f"escribir_matriz.py falló ({proc.returncode}): {proc.stderr[-2000:]}")
            return _agregar_porque((d / "salida.xlsx").read_bytes(), salida, libro)


HOJA_MATRIZ = "Extracción completa.csv"


def _agregar_porque(datos: bytes, salida: list[dict[str, Any]], libro: dict[str, Any]) -> bytes:
    """Suma a la nota de cada celda el porqué (inferencia o discrepancia del artículo).

    El script oficial solo escribe la evidencia en las notas; esto lo completa sin modificarlo.
    """
    posiciones = {k: c["posicion"] for k, c in libro["columnas"].items()}
    wb = openpyxl.load_workbook(io.BytesIO(datos))
    if HOJA_MATRIZ not in wb.sheetnames:
        return datos
    ws = wb[HOJA_MATRIZ]
    filas_por_estudio: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v is not None:
            filas_por_estudio[str(v)] = r
    for item in salida:
        fila = filas_por_estudio.get(str(item["fila"].get("Estudio")))
        if fila is None:
            continue
        tz = item["trazabilidad"]
        porque: dict[str, list[str]] = {}
        for x in tz.get("inferencias", []):
            porque.setdefault(x["columna"].split(" (")[0], []).append(
                f"POR QUÉ (inferido, {x['regla_aplicada']}): {x['razonamiento']}")
        for x in tz.get("discrepancias_en_el_articulo", []):
            porque.setdefault(x["columna"], []).append(
                f"DISCREPANCIA DEL ARTÍCULO: {x['fuente_1']} || {x['fuente_2']} -> adoptado: {x['valor_adoptado']}")
        for clave, lineas in porque.items():
            p = posiciones.get(clave)
            if p is None:
                continue
            celda = ws.cell(fila, p)
            previo = f"{celda.comment.text}\n\n" if celda.comment else ""
            nota = Comment((previo + "\n".join(lineas))[:3000], "Extractor")
            nota.width, nota.height = 420, 260
            celda.comment = nota
    salida_xlsx = io.BytesIO()
    wb.save(salida_xlsx)
    return salida_xlsx.getvalue()
