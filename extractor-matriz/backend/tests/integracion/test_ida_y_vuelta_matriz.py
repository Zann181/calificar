"""Aceptación de F1: importar y exportar la matriz de 71 filas da los mismos valores celda por celda."""

from __future__ import annotations

import io
import uuid

import openpyxl
import pytest

from app.composicion import Contenedor
from app.contextos.matriz.aplicacion.casos_de_uso import MatrizNoVacia
from app.contextos.matriz.dominio.fila import Origen
from tests.conftest import MATRIZ

HOJA = "Extracción completa.csv"


def test_importar_71_filas_con_25_notas(contenedor: Contenedor, proyecto: uuid.UUID) -> None:
    r = contenedor.matriz.importar.ejecutar(proyecto, MATRIZ.read_bytes(), MATRIZ.name)
    assert (r.filas, r.documentos, r.notas_heredadas) == (71, 47, 25)
    filas = contenedor.matriz.consultas.todas(proyecto)
    assert [f.estudio for f in filas] == list(range(1, 72))
    assert all(f.origen == Origen.HEREDADA for f in filas)
    kumar = contenedor.matriz.consultas.fila(proyecto, 1)
    assert kumar.valores()["Correlación Feli 1 - JP"] == 0.73
    assert sum(len(f.trazabilidad["notas_heredadas"]) for f in filas) == 25


def test_ida_y_vuelta_celda_por_celda(contenedor: Contenedor, proyecto: uuid.UUID) -> None:
    original = MATRIZ.read_bytes()
    contenedor.matriz.importar.ejecutar(proyecto, original, MATRIZ.name)
    exportacion = contenedor.exportar.ejecutar(proyecto)
    assert (exportacion.filas_heredadas, exportacion.filas_extraidas) == (71, 0)
    exportado = contenedor.exportar.leer(proyecto, exportacion.id)

    a = openpyxl.load_workbook(io.BytesIO(original))
    b = openpyxl.load_workbook(io.BytesIO(exportado))
    assert a.sheetnames == b.sheetnames
    wa, wb = a[HOJA], b[HOJA]
    columnas = contenedor.normas.consultas.columnas(proyecto)
    diferencias = []
    celdas = 0
    for fila in range(1, wa.max_row + 1):
        for c in columnas:
            x, y = wa.cell(fila, c.posicion), wb.cell(fila, c.posicion)
            celdas += 1
            if x.value != y.value:
                diferencias.append((fila, c.clave, x.value, y.value))
            if (x.comment and x.comment.text) != (y.comment and y.comment.text):
                diferencias.append((fila, c.clave, "comentario"))
    assert celdas >= 72 * 54
    assert diferencias == []


def test_reimportar_exige_reemplazar(contenedor: Contenedor, proyecto: uuid.UUID) -> None:
    datos = MATRIZ.read_bytes()
    contenedor.matriz.importar.ejecutar(proyecto, datos, MATRIZ.name)
    with pytest.raises(MatrizNoVacia):
        contenedor.matriz.importar.ejecutar(proyecto, datos, MATRIZ.name)
    r = contenedor.matriz.importar.ejecutar(proyecto, datos, MATRIZ.name, reemplazar=True)
    assert r.reemplazadas == 71
    assert len(contenedor.matriz.consultas.todas(proyecto)) == 71


def test_exportar_sin_importar(contenedor: Contenedor, proyecto: uuid.UUID) -> None:
    with pytest.raises(LookupError, match="importe la matriz"):
        contenedor.exportar.ejecutar(proyecto)
